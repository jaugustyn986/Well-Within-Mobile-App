from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Well_Within_Clinical_Review_Packet.docx"


# compact_reference_guide preset with one named brand override:
# blue accents -> Well Within muted rose/charcoal; all geometry and rhythm remain preset-based.
FONT = "Calibri"
INK = "3E3533"
MUTED = "6F6663"
ACCENT = "A96F68"
ACCENT_DARK = "754B47"
PALE_ROSE = "F5EAE7"
PALE_SAGE = "EAF1ED"
PALE_GOLD = "F6F0DD"
LIGHT_GRAY = "F3F1F0"
WHITE = "FFFFFF"
BORDER = "D8CFCC"
LINK = "7C4F70"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 100, "bottom": 100, "start": 120, "end": 120}


@dataclass(frozen=True)
class Decision:
    id: str
    title: str
    audit_ids: str
    why: str
    current: str
    evidence: str
    questions: Sequence[str]


DECISIONS: tuple[Decision, ...] = (
    Decision(
        "D-01",
        "Daily observation reduction",
        "R-01; I-08",
        "When a user records more than one observation on the same day, the app must decide what is retained and what can drive a derived marker.",
        "The engine normalizes each observation, then uses the highest internal 0-3 rank as the day's interpretive value. The original observations remain stored.",
        "Public materials support recording the strongest sign in broad terms. The app's rank abstraction, tie handling, and invalid-combination behavior do not yet have source-level approval.",
        (
            "Is strongest-sign reduction appropriate for every supported user and cycle context?",
            "When two observations conflict, should the app select, preserve both without interpretation, or ask the user to correct the entry?",
            "Which original details must remain visible to a reviewer even when one value drives the automated summary?",
        ),
    ),
    Decision(
        "D-02",
        "Sensation, appearance, and generated notation",
        "R-02-R-05; R-07; I-01-I-03; I-09",
        "The mapping from recorded words to generated notation and internal ranks determines every later summary.",
        "The app maps sensations, appearances, and frequency into generated notation and a 0-3 interpretive rank. Damp, wet, and shiny inputs can currently open the derived mucus pattern.",
        "The repository contains an unsourced internal table and a direct contradiction about whether several lower-numbered observations represent mucus or non-mucus sensations. Authorized source edition and page-level review are missing.",
        (
            "Using an authorized source, which mappings are clinically correct for software entry and export?",
            "Should damp, wet, or shiny without a mucus appearance open any automated pattern boundary?",
            "Does omitted frequency affect interpretation, export notation only, or entry validity?",
        ),
    ),
    Decision(
        "D-03",
        "Contradictory or incomplete input combinations",
        "R-06; I-04-I-06; I-10",
        "A polished interface can still accept combinations that a trained practitioner would regard as internally inconsistent.",
        "Some appearance selections can override a 'dry' sensation and become Peak-type in the current engine. Missing dates are tolerated in legacy/test paths, while production records normally carry dates.",
        "No documented clinical decision defines which combinations are valid, which require correction, or which may be recorded but excluded from interpretation.",
        (
            "For dry plus clear, lubricative, or yellow appearance, should the app accept, warn, require correction, or record without interpretation?",
            "Which fields are mandatory before a day may participate in automated interpretation?",
            "What recovery message is appropriate for invalid, duplicate, missing, or out-of-order dates?",
        ),
    ),
    Decision(
        "D-04",
        "Flow, spotting, brown observations, and mucus",
        "R-09-R-12; F-06; I-07; C-10-C-11",
        "Users can record mucus during bleeding, but the current interpretation treatment differs by bleeding class.",
        "Menstrual flow suppresses Peak-type interpretation on that day. Spotting remains descriptive. Light/spotting rows are currently blocked from opening the derived interval or becoming a Peak candidate.",
        "Public method material indicates mucus is observed on light/very-light flow days, but the correct automated treatment and wording have not been approved.",
        (
            "Can a light-flow or spotting day influence a Peak-type marker, and under what exact conditions?",
            "Which bleeding observations are descriptive only, and which require individualized review?",
            "What should the ordinary user see when bleeding and a high-quality mucus sign occur together?",
        ),
    ),
    Decision(
        "D-05",
        "Cycle start and continuing flow",
        "R-21-R-22; F-07; C-12-C-14",
        "Every cycle day, Peak day, duration, comparison, and export depends on the cycle-start boundary.",
        "The current engine begins a cycle at a heavy/moderate flow row that is not treated as continuing prior flow. Calendar-day math is now used for displayed day numbers and durations, but boundary adjudication remains unresolved.",
        "Public descriptions refer broadly to onset of menstruation; the exact heavy/moderate threshold, leading light days, missing dates within flow, and charts without heavy/moderate flow lack qualified approval.",
        (
            "For light -> heavy -> moderate, which date is Cycle Day 1 and which dates belong to the cycle?",
            "How should a missing calendar date between heavy/moderate observations affect continuing-flow logic?",
            "What should the app do when no heavy/moderate boundary exists?",
        ),
    ),
    Decision(
        "D-06",
        "Peak-type observation definition",
        "R-13; C-13; I-01-I-08",
        "The app must separate what the user recorded from the later retrospective Peak Day marker.",
        "Clear, cloudy/clear, lubricative, stretchy, and some promoted combinations currently produce the highest internal rank and may be labeled Peak-type based on what was recorded.",
        "Clear, stretchy, and lubricative signs are broadly supported in public and peer-reviewed descriptions. Exact combination handling and the cloudy/clear mapping require authorized review.",
        (
            "Which exact observations may be called 'Peak-type' in an unassisted app?",
            "Is 'Peak-type observation based on what you recorded' acceptable ordinary-user wording?",
            "Which observations must never be promoted to Peak-type automatically?",
        ),
    ),
    Decision(
        "D-07",
        "Retrospective Peak Day and the three-day sequence",
        "R-14-R-16; C-01; C-05; C-15",
        "This is the central derived marker currently surfaced throughout Calendar and Cycle History.",
        "A non-flow highest-rank day becomes a candidate. The app marks it as Peak Day only after the next three consecutive calendar days are present, observed, and each lower than the candidate. A later highest-rank day during the waiting period resets the candidate.",
        "The last clear/stretchy/lubricative day and P+3 framework are publicly described. The strict software comparison, reset behavior, and permitted patient-facing language have not been clinically approved for this implementation.",
        (
            "Is the strict 'each of the next three days is lower than the candidate' rule correct for the proposed scope?",
            "Does a gradual decline require any relationship among P+1, P+2, and P+3 beyond all being lower than the candidate?",
            "Which terms are acceptable: candidate, identified Peak Day, marked Peak Day, P+1/P+2/P+3, post-Peak pattern?",
        ),
    ),
    Decision(
        "D-08",
        "Calendar gaps and days marked not observed",
        "R-25; C-02-C-04",
        "The app should be conservative without trapping users or misrepresenting an intentionally unobserved day as an app error.",
        "A missing calendar date or explicit not-observed day inside the three-day confirmation sequence blocks the Peak summary. An earlier gap can limit the pattern boundary. Future days beyond the last record are treated as a developing pattern, not a gap. Charting and editing always remain available.",
        "This behavior is internally verified and deliberately conservative; exact clinical effect and recovery wording still need review.",
        (
            "Should an absent date and an explicit not-observed row have the same clinical effect?",
            "Which gaps limit only the opening boundary, which block Peak, and which make the full pattern unsupported?",
            "Is the proposed guidance appropriate: keep charting, restore a remembered open day, and preserve an intentionally not-observed day?",
        ),
    ),
    Decision(
        "D-09",
        "Separated or repeated Peak-type sequences",
        "R-17; F-05; C-09",
        "Selecting the first or last sequence without an approved rule could create a clinically consequential false conclusion.",
        "When more than one separated sequence independently satisfies the current three-lower-days rule, Well Within shows 'Your chart shows more than one possible Peak pattern' and does not choose one Peak Day. Outside review is optional and never unlocks the app.",
        "Repeated/double-Peak behavior is outside the implemented method scope. The current review state is a conservative product decision, not a clinical classification.",
        (
            "Is non-selection plus optional review the correct default for two independently qualifying sequences?",
            "Are there any patterns where an automated app may select the later sequence? State the exact prerequisites.",
            "What wording avoids diagnosis while accurately explaining why one Peak Day is not shown?",
        ),
    ),
    Decision(
        "D-10",
        "Continuous mucus, non-Peak-only patterns, and special contexts",
        "R-19-R-20; C-06-C-08; special-context set",
        "A standard dry-to-mucus-to-Peak model must not silently imply that every chart follows that pattern.",
        "Continuous Peak-type signs remain a developing pattern. Continuous lower-quality mucus and non-Peak-only runs do not receive a confirmed Peak summary. No BIP, postpartum, perimenopause, medication, infection, or diagnostic mechanics have been invented.",
        "Public materials describe BIP and special-context complexity, but this app has no approved interpretation path or population scope for those patterns.",
        (
            "For each special context, choose: full interpretation, restricted interpretation, defined caution, or no interpretation/referral.",
            "What minimum history or practitioner-established BIP is required before any automated summary?",
            "Which neutral message should appear when the app intentionally does not interpret the pattern?",
        ),
    ),
    Decision(
        "D-11",
        "Permitted ordinary-user terminology",
        "C-01-C-19; R-18; R-31-R-32",
        "Users need understandable language without converting a chart marker into ovulation, biological certainty, or contraceptive guidance.",
        "The interim UI uses observational language: possible Peak Day, chart marks a Peak Day after three lower days, post-Peak pattern, missing-data limitation, or multiple possible Peak patterns. Confidence tiers, 'most fertile day confirmed,' and broad current fertile-window conclusions were removed or suppressed.",
        "The wording is product-reviewed and internally consistent, but it is not clinically or legally approved. 'Fertile Start,' 'Fertile End,' 'total fertile days,' and equivalent claims remain held.",
        (
            "Approve, revise, or prohibit each proposed term in the terminology worksheet.",
            "Specify the exact limitation that must appear next to any retrospective Peak Day statement.",
            "Identify any wording that could improperly support avoiding-pregnancy or diagnosis decisions.",
        ),
    ),
    Decision(
        "D-12",
        "History comparisons and inclusion eligibility",
        "R-29-R-30; C-21-C-22",
        "Retrospective comparisons can feel predictive or normative if thresholds, sample size, and eligible-cycle rules are unclear.",
        "Only cycles with the summary-available support state contribute to derived aggregates, comparison context, and Peak overlays. Active-cycle anticipatory timing is suppressed. Calendar-correct dates are used. Existing comparison bands remain product thresholds without clinical meaning.",
        "No clinical source approves the current +/- day thresholds or terms such as usual, consistent, or significant variation.",
        (
            "Which comparisons are purely descriptive and appropriate for ordinary users?",
            "What minimum eligible sample size and exclusions are required?",
            "Should comparison bands be removed, numerically displayed without labels, or clinically approved against evidence?",
        ),
    ),
    Decision(
        "D-13",
        "Possible fertile pattern presentation",
        "R-33-R-35; C-25-C-27; PFP-01-PFP-06",
        "Users may benefit from seeing what a fertile pattern could be from their recorded data, but an exact band can be mistaken for biological certainty, prediction, or pregnancy-avoidance guidance.",
        "The app already computes an opening index and P+3 end index and renders a Fertile Window timeline. Phase 1C proposes reusing that foundation with the qualified label 'Possible fertile pattern,' no exact dates while forming, exact dates only for an eligible summary-available chart, and suppression for blocked/review/unsupported charts. This direction is documented but not implemented.",
        "The public Saint Paul VI sample describes beginning of mucus through three full days past Peak among days of fertility for achieving pregnancy, but it also lists additional bleeding and non-Peak rules and requires trained instruction. No public evidence validates Well Within's exact opening, Peak/P+3, eligibility, wording, or user comprehension.",
        (
            "For an eligible simple chart, is 'Possible fertile pattern — based on your logged observations' acceptable from the first approved mucus observation through P+3?",
            "While a chart is forming, should the app state only that a possible pattern may be developing, without start/end dates or a band?",
            "Confirm the exact suppression, later-sign, history-range, and adjacent-limitation behavior in PFP-01 through PFP-06.",
        ),
    ),
)


CORE_FIXTURES = (
    ("C-01", "Standard dry -> mucus -> Peak -> P+3", "Flow; dry; lower-quality mucus; wet; clear/stretchy/lubricative-type day; three lower observed calendar days", "Mark cycle start, pattern boundary, Peak-type days, Peak Day, P+1-P+3, allowed wording, and history eligibility."),
    ("C-02", "Gap before pattern opening", "Flow; dry; one absent calendar date; first mucus observation", "Does the gap block, qualify, or leave the opening boundary unchanged?"),
    ("C-03", "Gap during three-day sequence", "Peak-type day; one lower observed day; absent date; later lower row", "Confirm blocked state, affected outputs, and recovery wording."),
    ("C-04", "Explicit not observed", "Peak-type day; one day intentionally marked not observed; two lower rows", "Distinguish intentionally not observed from an absent date, if clinically necessary."),
    ("C-05", "Gradual decline", "Peak-type day; wet; damp; dry", "Confirm whether all three days being lower is sufficient and identify the Peak date."),
    ("C-06", "Continuous Peak-type signs", "Clear/stretchy/lubricative-type observations for 10-20 consecutive days", "Choose supported, restricted, caution, or no interpretation; provide safe wording."),
    ("C-07", "Continuous non-Peak mucus / BIP-like", "Unchanged sticky/cloudy observations for 30 days", "Define the prerequisite for any interpretation and whether standard summaries must remain suppressed."),
    ("C-08", "Non-Peak-only run ending", "One to two lower-quality mucus days then dry; separately three or more then dry", "Define whether any Peak/plus-count/end concept is permitted or the pattern remains unsupported."),
    ("C-09", "Separated qualifying sequences", "Peak-type; three lower days; later Peak-type; three lower days", "Choose first, later, review without selection, or no interpretation."),
    ("C-10", "High-quality mucus during light flow", "Heavy; moderate; light plus clear/lubricative sign; three lower days", "Can the flow day be a candidate, be recorded only, or require review?"),
    ("C-11", "Mid-cycle spotting or brown observation", "Isolated spotting/brown observation followed by dry days", "Define recording, boundary, fertility-related meaning, and review wording."),
    ("C-12", "Cycle-start preamble", "Light; heavy; moderate", "Identify Cycle Day 1 and which days belong in cycle length."),
    ("C-13", "Missing date inside flow", "Heavy/moderate Jan 1; no Jan 2 row; heavy/moderate Jan 3", "Continuing flow, possible new cycle, or uninterpretable?"),
    ("C-14", "No heavy/moderate boundary", "Complete record containing only light/spotting/brown bleeding", "Create a cycle, leave unbounded, or require review?"),
    ("C-15", "Past edit changes marker", "Qualifying sequence; later edit changes a post-candidate day to Peak-type", "Confirm recalculation, removal/replacement of the prior marker, and user message."),
)


INPUT_FIXTURES = (
    ("I-01", "Damp with no appearance", "Mucus, non-mucus sensation, incomplete entry, or other?"),
    ("I-02", "Wet with no appearance", "Classification and entry-validity decision."),
    ("I-03", "Shiny with no appearance", "Classification and entry-validity decision."),
    ("I-04", "Dry plus clear appearance", "Accept, correct, record without interpretation, or Peak-type?"),
    ("I-05", "Dry plus lubricative appearance", "Accept, correct, record without interpretation, or Peak-type?"),
    ("I-06", "Dry plus yellow appearance", "Valid combination and classification?"),
    ("I-07", "Spotting plus clear/lubricative sign", "Primary display and whether it may affect Peak."),
    ("I-08", "Several observations with conflicting signs", "Confirm reduction logic and details that must remain visible."),
    ("I-09", "Frequency omitted vs. recorded", "Interpretive effect, export-only effect, or required correction?"),
    ("I-10", "Duplicate, out-of-order, or invalid dates", "Required rejection, recovery, and reviewer-facing representation."),
)


TERMS = (
    ("Peak-type observation", "Current interim use", "Approve / revise / prohibit", "Describes a recorded sign, not ovulation."),
    ("Possible Peak Day", "Current interim use", "Approve / revise / prohibit", "Used while the three following days are incomplete."),
    ("Well Within marked Cycle Day X as Peak Day", "Current interim use", "Approve / revise / prohibit", "Retrospective software marker after the present rule is met."),
    ("Your chart shows a post-Peak pattern", "Current interim use", "Approve / revise / prohibit", "Chart-only statement; no ovulation claim."),
    ("This reflects your chart; it does not confirm ovulation", "Current limitation", "Approve / revise / require", "Shown beside supported retrospective Peak summaries."),
    ("A few days need context", "Current interim use", "Approve / revise / prohibit", "Missing/absent date limits what the app shows."),
    ("More than one possible Peak pattern", "Current interim use", "Approve / revise / prohibit", "App does not choose between separated qualifying sequences."),
    ("Fertile Start / Fertile End / Total fertile days", "Held", "Approve only with defined scope / continue hold", "Broad patient-specific boundary terms."),
    ("High / Moderate / Low confidence", "Removed", "Keep removed / define validated alternative", "No calibrated probability supports these labels."),
    ("Most fertile day confirmed / past the fertile window", "Held/removed", "Continue hold / provide approved alternative", "Can imply biological certainty or decision guidance."),
    ("Possible fertile pattern - based on your logged observations", "Phase 1C working direction", "Approve / revise / prohibit for eligible scope", "Qualified chart-based estimate; not a complete CrMS determination or prediction."),
    ("This reflects your charted observations. It does not confirm ovulation, identify infertile days, or provide pregnancy-avoidance guidance.", "Phase 1C proposed limitation", "Approve / revise / require", "Shown beside any exact possible-pattern dates."),
)


SPECIAL_CONTEXTS = (
    "Postpartum and breastfeeding",
    "Perimenopause",
    "Recent hormonal contraception",
    "Persistent discharge or suspected infection",
    "Irregular or unusual bleeding",
    "Medications or conditions that may alter signs",
    "Very long cycles with several mucus patches",
    "Known subfertility",
    "Fewer than two eligible historical cycles",
)


METHOD_SOURCES = (
    ("Creighton Model - About", "https://creightonmodel.com/about/", "Standardized, practitioner-taught system and individualized follow-up; does not validate this app."),
    ("Creighton Model - Background", "https://creightonmodel.com/background/", "Public overview of mucus progression, Peak, long-cycle patches, and BIP complexity; not a software license."),
    ("Saint Paul VI Institute public chart/IP sample", "https://saintpaulvi.com/PDF/CrMS_App_Copyright.pdf", "Public sample includes a beginning-of-mucus-through-P+3 concept plus additional rules and protected-material warnings; not a complete specification, license, or validation of Well Within."),
    ("FertilityCare Centers of America - Creighton Model System", "https://www.fertilitycare.org/creighton-model-system/", "Current description of a standardized, practitioner-taught system; no approval or affiliation with Well Within."),
    ("Peak Day selected by woman, expert, and computer", "https://pubmed.ncbi.nlm.nih.gov/32101336/", "Shows that Peak identification can differ; does not validate Well Within's algorithm."),
    ("Cervical mucus patterns and the fertile window", "https://pmc.ncbi.nlm.nih.gov/articles/PMC8487651/", "Trained-participant observations and non-Peak-only pattern nuance; not universal app validation."),
    ("Peak Day validation against a hormone monitor", "https://pmc.ncbi.nlm.nih.gov/articles/PMC3893397/", "Trained-instruction context and limits; does not approve Well Within wording or classifications."),
    ("CDC U.S. MEC - Fertility Awareness-Based Methods", "https://www.cdc.gov/contraception/hcp/usmec/fertility-awareness-based-methods.html", "Identifies contexts requiring delay or special counseling; does not define this app's method rules."),
)


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_margins(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        if edge in kwargs:
            tag = "w:" + edge
            node = tc_mar.find(qn(tag))
            if node is None:
                node = OxmlElement(tag)
                tc_mar.append(node)
            node.set(qn("w:w"), str(kwargs[edge]))
            node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    """Keep a review or worksheet row intact across page boundaries."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table, color: str = BORDER, size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: Sequence[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, **CELL_MARGINS_DXA)


def set_font(run, *, size: float | None = None, color: str = INK, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.color.rgb = rgb(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_keep(paragraph, *, keep_next: bool = False, keep_lines: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    if keep_lines:
        keep = OxmlElement("w:keepLines")
        p_pr.append(keep)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.extend((r_fonts, color, underline))
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, value, end))


def add_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element

    def next_id(tag: str, attr: str) -> int:
        vals = [int(x.get(qn(attr))) for x in numbering.findall(qn(tag)) if x.get(qn(attr))]
        return max(vals, default=0) + 1

    def create(fmt: str, text: str, marker_font: str | None = None) -> int:
        abstract_id = next_id("w:abstractNum", "w:abstractNumId")
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.extend((tabs, ind))
        lvl.extend((start, num_fmt, lvl_text, suff, lvl_jc, p_pr))
        if marker_font:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), marker_font)
            fonts.set(qn("w:hAnsi"), marker_font)
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

        num_id = next_id("w:num", "w:numId")
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        return num_id

    return create("bullet", "•", "Symbol"), create("decimal", "%1.")


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num))
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def add_bullet(doc: Document, text: str, bullet_num_id: int) -> None:
    p = doc.add_paragraph()
    apply_num(p, bullet_num_id)
    run = p.add_run(text)
    set_font(run, size=11)


def add_numbered(doc: Document, text: str, decimal_num_id: int) -> None:
    p = doc.add_paragraph()
    apply_num(p, decimal_num_id)
    run = p.add_run(text)
    set_font(run, size=11)


def add_body(doc: Document, text: str, *, bold_lead: str | None = None, italic: bool = False) -> None:
    p = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, size=11, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest, size=11, italic=italic)
    else:
        run = p.add_run(text)
        set_font(run, size=11, italic=italic)


def add_callout(doc: Document, label: str, text: str, fill: str = PALE_ROSE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [PAGE_WIDTH_DXA])
    set_table_borders(table, color=fill, size="2")
    set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    label_run = p.add_run(label + "  ")
    set_font(label_run, size=10.5, color=ACCENT_DARK, bold=True)
    text_run = p.add_run(text)
    set_font(text_run, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_paragraph_keep(p, keep_next=True)


def format_table_text(table, *, header: bool = False, size: float = 9.0) -> None:
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.10
                for run in p.runs:
                    set_font(run, size=size, bold=(header and row_index == 0))


def add_label_detail_table(doc: Document, rows: Sequence[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [1700, 7660])
    set_table_borders(table)
    for idx, (label, value) in enumerate(rows):
        row = table.rows[idx]
        set_row_cant_split(row)
        left, right = row.cells
        set_cell_shading(left, LIGHT_GRAY)
        left.text = label
        right.text = value
        for run in left.paragraphs[0].runs:
            set_font(run, size=9.2, color=ACCENT_DARK, bold=True)
        for run in right.paragraphs[0].runs:
            set_font(run, size=9.2)
    format_table_text(table, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_response_block(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [PAGE_WIDTH_DXA])
    set_table_borders(table, color=BORDER)
    cell = table.cell(0, 0)
    set_row_cant_split(table.rows[0])
    set_cell_shading(cell, PALE_SAGE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Reviewer decision")
    set_font(r, size=10, color=ACCENT_DARK, bold=True)
    for line in (
        "[ ] Approve as implemented for the stated scope",
        "[ ] Approve with the changes written below",
        "[ ] Do not interpret automatically; use a defined review/support state",
        "[ ] Hold pending an authorized source or additional evidence",
    ):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        set_font(r, size=9.5)
    for label in ("Required change / limitation:", "Authorized source, edition, and page:", "Reviewer initials and date:"):
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(label + " " + "_" * 62)
        set_font(r, size=9.2, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_fixture_table(doc: Document, rows: Sequence[tuple[str, ...]], headers: Sequence[str], widths: Sequence[int], size: float = 8.4) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0]
    for cell, text in zip(hdr.cells, headers):
        cell.text = text
        set_cell_shading(cell, PALE_ROSE)
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for row_values in rows:
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        for cell, text in zip(cells, row_values):
            cell.text = text
    set_table_geometry(table, widths)
    set_table_borders(table)
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1.5)
                p.paragraph_format.line_spacing = 1.04
                if col_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_font(run, size=size, bold=(row_idx == 0), color=(ACCENT_DARK if row_idx == 0 else INK))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = rgb(ACCENT_DARK)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.0

    subtitle = styles["Subtitle"]
    subtitle.font.name = FONT
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    subtitle.font.size = Pt(13.5)
    subtitle.font.color.rgb = rgb(MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.line_spacing = 1.15

    heading_tokens = {
        1: (16, ACCENT_DARK, 18, 10),
        2: (13, ACCENT, 14, 7),
        3: (12, ACCENT_DARK, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Small Note" not in styles:
        note = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = styles["Small Note"]
    note.font.name = FONT
    note._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    note._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    note.font.size = Pt(9)
    note.font.color.rgb = rgb(MUTED)
    note.paragraph_format.space_after = Pt(4)
    note.paragraph_format.line_spacing = 1.10


def configure_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("WELL WITHIN  |  CLINICAL REVIEW PACKET")
    set_font(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [7200, 2160], indent_dxa=0)
    left, right = table.rows[0].cells
    left_p = left.paragraphs[0]
    left_p.paragraph_format.space_after = Pt(0)
    r = left_p.add_run("Reviewer working document | Version 1.1 | July 13, 2026")
    set_font(r, size=8, color=MUTED)
    add_page_field(right.paragraphs[0])

    first_footer = section.first_page_footer
    p = first_footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Well Within | Clinical review requested; not a clinical validation or legal approval")
    set_font(r, size=8, color=MUTED)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("QUALIFIED REVIEWER PACKET")
    set_font(r, size=10.5, color=ACCENT, bold=True)

    title = doc.add_paragraph("Well Within", style="Title")
    set_paragraph_keep(title, keep_next=True)
    subtitle = doc.add_paragraph("Clinical Rules, Pattern Boundaries, and Fixture Review", style="Subtitle")
    set_paragraph_keep(subtitle, keep_next=True)

    meta = doc.add_table(rows=4, cols=2)
    meta_rows = (
        ("Prepared for", "A qualified current Creighton/FABM clinical reviewer"),
        ("Review purpose", "Adjudicate exact software rules, supported scope, and permitted terminology"),
        ("Implementation", "Well Within rules engine at commit 261fc0f"),
        ("Packet version", "1.1 - July 13, 2026"),
    )
    for row, (label, value) in zip(meta.rows, meta_rows):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        for run in row.cells[0].paragraphs[0].runs:
            set_font(run, size=9.5, color=ACCENT_DARK, bold=True)
        for run in row.cells[1].paragraphs[0].runs:
            set_font(run, size=9.5)
    set_table_geometry(meta, [1900, 7460])
    set_table_borders(meta)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_callout(
        doc,
        "What we are asking",
        "Please review the stated current behavior against authorized sources and your clinical judgment. For each decision, approve the behavior for a clearly defined scope, specify a change, require a review/support state, or hold it pending evidence.",
        PALE_SAGE,
    )
    add_callout(
        doc,
        "What we are not asking",
        "This is not a request for blanket endorsement, certification, affiliation, diagnosis, ovulation confirmation, pregnancy prediction, or avoiding-pregnancy guidance. Legal, regulatory, privacy, trademark, copyright, and licensing decisions are routed separately to counsel.",
        PALE_GOLD,
    )

    p = doc.add_paragraph(style="Small Note")
    r = p.add_run("Confidentiality note: This packet contains product behavior and synthetic review scenarios only. It includes no private user charts or personal health information.")
    set_font(r, size=9, color=MUTED, italic=True)


def add_orientation(doc: Document, bullet_id: int, number_id: int) -> None:
    doc.add_page_break()
    add_heading(doc, "1. Reviewer orientation", 1)
    add_body(doc, "Well Within is a local-first mobile charting app. Users record daily bleeding and mucus observations. A deterministic, network-free rules engine organizes those observations into chart markers and retrospective summaries. The app does not claim to diagnose a condition, confirm ovulation, predict pregnancy, or replace individualized instruction.")

    add_heading(doc, "The narrow review objective", 2)
    for text in (
        "Determine which exact observation mappings and transformations are clinically appropriate for an unassisted app.",
        "Define the population and chart patterns for which any automated retrospective marker is supportable.",
        "Decide when the app should continue showing observations without an automated conclusion.",
        "Approve, revise, or prohibit the ordinary-user terms listed in this packet.",
        "Identify the authorized source edition and page for every method-specific rule that is approved.",
    ):
        add_bullet(doc, text, bullet_id)

    add_heading(doc, "Requested review method", 2)
    for text in (
        "Review each priority decision independently; do not infer approval from the current implementation.",
        "Mark one response category and write any required change, limitation, or source citation.",
        "Review the fixture set and annotate the expected gold output for each applicable scenario.",
        "Provide at least 20 de-identified, authorized gold charts spanning supported and unsupported cases before final production approval.",
        "Join one adjudication session with product and engineering to resolve conflicts and version the approved decision record.",
    ):
        add_numbered(doc, text, number_id)

    add_heading(doc, "Response categories", 2)
    for label, detail in (
        ("Approve — ", "Appropriate for the stated population and pattern scope."),
        ("Approve with changes — ", "Use only after the reviewer’s written change is made."),
        ("Review/support state — ", "Keep observations visible; do not choose a marker automatically."),
        ("Hold — ", "No automated output until authorized evidence is supplied."),
    ):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.08
        lead = p.add_run(label)
        set_font(lead, size=10, bold=True)
        rest = p.add_run(detail)
        set_font(rest, size=10)

    add_callout(doc, "Evidence standard", "Repository specifications and automated tests show implementation consistency. They do not substitute for authorized method sources, qualified clinical judgment, or product-specific validation.")


def add_current_model(doc: Document, bullet_id: int) -> None:
    add_heading(doc, "2. Current implementation and interim safeguards", 1)
    add_body(doc, "The following describes what the current product does. It is not presented as a clinically approved model.")

    add_heading(doc, "Processing sequence", 2)
    for text in (
        "Observation: the user records bleeding, sensation, appearance, frequency, notes, or an intentionally not-observed day.",
        "Normalization: the engine derives a daily mucus classification, internal rank, bleeding class, and optional generated notation.",
        "Marker logic: the engine evaluates a possible Peak-type day and a three-lower-calendar-days sequence.",
        "Capability state: the app decides whether a retrospective summary is available, still forming, limited by missing data, or not selected because more than one sequence qualifies.",
        "Presentation: Calendar, Cycle History, Help, and PDF exports consume the same support state and suppress ineligible derived conclusions.",
    ):
        add_bullet(doc, text, bullet_id)

    add_heading(doc, "Current capability states", 2)
    add_fixture_table(
        doc,
        (
            ("forming", "The record does not yet support a retrospective summary.", "Names the observations or developing pattern and asks the user to keep charting."),
            ("summary available", "One sequence satisfies the current software rule.", "Names the marked Peak Day, the three following days, and the chart-only limitation."),
            ("blocked by missing", "A gap or not-observed day affects the current boundary or confirmation sequence.", "Explains the specific limitation; charting and editing remain available."),
            ("review recommended", "More than one separated sequence independently satisfies the current rule.", "Does not select one Peak Day; outside support is optional and never unlocks the app."),
        ),
        ("State", "Current trigger", "Current user path"),
        (1800, 3300, 4260),
        size=8.8,
    )

    add_heading(doc, "Engineering changes already completed", 2)
    for text in (
        "Displayed cycle days, Peak days, durations, comparisons, overlays, and PDF rows now use elapsed calendar dates rather than stored-row indexes.",
        "The unsupported historical shortcut that estimated fertile opening as Peak minus five days was removed.",
        "High/Moderate/Low confidence language was replaced with capability and completeness language.",
        "Missing/review states suppress derived summaries across Cycle Detail, History, comparisons, overlays, and default PDF output while preserving observations and export access.",
        "Charting, editing, history, and observation export remain available in every support state.",
    ):
        add_bullet(doc, text, bullet_id)

    add_callout(doc, "Verification context", "At packet creation, 184 core rules tests and 53 mobile tests pass; CI coverage is 97% statements, 92.33% branches, 100% functions, and 98.66% lines. This proves internal consistency only, not clinical validity.", PALE_SAGE)


def add_decisions(doc: Document, bullet_id: int) -> None:
    add_heading(doc, "3. Priority clinical decision records", 1)
    add_body(doc, "Complete every record. If approval depends on a different population, pattern, training relationship, or source edition, write that condition explicitly.")
    add_heading(doc, "Decision map", 2)
    add_fixture_table(
        doc,
        tuple((decision.id, decision.title) for decision in DECISIONS),
        ("ID", "Decision"),
        (1200, 8160),
        size=9.1,
    )
    for index, decision in enumerate(DECISIONS):
        # The decision map already fills the preceding page. Let the first record
        # flow naturally to the next page so Word does not introduce a blank page.
        if index > 0:
            doc.add_page_break()
        add_heading(doc, f"{decision.id}  {decision.title}", 2)
        add_label_detail_table(
            doc,
            (
                ("Audit links", decision.audit_ids),
                ("Why it matters", decision.why),
                ("Current behavior", decision.current),
                ("Evidence state", decision.evidence),
            ),
        )
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("Questions for the reviewer")
        set_font(r, size=10, color=ACCENT_DARK, bold=True)
        set_paragraph_keep(p, keep_next=True)
        for question in decision.questions:
            add_bullet(doc, question, bullet_id)
        add_response_block(doc)


def add_fixtures(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "4. Clinical fixture worksheet", 1)
    add_body(doc, "For every applicable fixture, independently mark the valid daily inputs, cycle start, cycle-day numbers, relevant mucus boundary, Peak-type observations, Peak candidate, retrospectively identified Peak Day, P+1-P+3, any permitted interval end, support state, allowed wording, and history eligibility.")
    add_callout(doc, "Important", "These are synthetic review scenarios, not private user charts. A final acceptance set still requires at least 20 de-identified, licensed or authorized gold charts reviewed against the exact observation fields.", PALE_GOLD)

    add_heading(doc, "Core scenarios", 2)
    add_fixture_table(
        doc,
        CORE_FIXTURES,
        ("ID", "Scenario", "Minimum sequence", "Reviewer output requested"),
        (620, 2000, 3200, 3540),
        size=8.0,
    )

    add_heading(doc, "Input-combination scenarios", 2)
    add_fixture_table(
        doc,
        INPUT_FIXTURES,
        ("ID", "Combination", "Reviewer decision requested"),
        (700, 3100, 5560),
        size=8.6,
    )

    add_heading(doc, "Special contexts requiring an explicit scope decision", 2)
    add_body(doc, "For each context, select one: [ ] full interpretation  [ ] restricted interpretation  [ ] defined caution  [ ] no interpretation/referral. Then describe prerequisites and permitted wording.")
    for context in SPECIAL_CONTEXTS:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run("[ ] " + context + "  -  Decision/limits: " + "_" * 48)
        set_font(r, size=10)

    add_heading(doc, "Gold-chart acceptance record", 2)
    for text in (
        "Reviewer marks each gold chart independently before seeing automated output.",
        "Product records exact allowed wording and supported population/pattern scope.",
        "Engineering encodes exact ISO dates and original observation fields, not only synthetic ranks.",
        "Output is compared field by field; disagreements are adjudicated and logged.",
        "Approval records reviewer identity/credentials, authorized source edition, date, engine version, and limitations.",
    ):
        add_body(doc, "[ ] " + text)
    add_heading(doc, "Gold-chart set details", 3)
    for label in (
        "Set identifier / custodian:",
        "Number of charts reviewed:",
        "Authorization or license reference:",
        "Adjudication record location:",
        "Exceptions or unresolved charts:",
    ):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run(label + " " + "_" * 68)
        set_font(r, size=10)


def add_terminology(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "5. Ordinary-user terminology worksheet", 1)
    add_body(doc, "Review the net impression of each term in the context of an unassisted health app. Approval should specify the exact text, required limitation, allowed surface, and supported scope.")
    add_fixture_table(
        doc,
        TERMS,
        ("Term", "Interim status", "Reviewer action", "Context / concern"),
        (2600, 1550, 2150, 3060),
        size=8.2,
    )
    add_callout(doc, "Non-negotiable interim boundary", "Until separately approved, Well Within will not present ovulation confirmation, pregnancy probability, diagnosis, treatment, 'safe'/'infertile' claims, or avoiding-pregnancy instructions.", PALE_GOLD)

    add_heading(doc, "Approved wording record", 2)
    for label in (
        "Observation-level term:",
        "Developing-pattern term:",
        "Retrospective Peak Day statement:",
        "Required chart-only / ovulation limitation:",
        "Missing-data statement:",
        "Multiple-sequence statement:",
        "Unsupported-context statement:",
        "Terms prohibited for ordinary users:",
    ):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(label + " " + "_" * 72)
        set_font(r, size=10)


def add_sources(doc: Document, number_id: int) -> None:
    doc.add_page_break()
    add_heading(doc, "6. Source and evidence register", 1)
    add_body(doc, "Please replace broad web support with the exact authorized source edition and page whenever approving a method-specific rule. The sources below provide context; none independently validates Well Within.")
    for title, url, note in METHOD_SOURCES:
        p = doc.add_paragraph()
        apply_num(p, number_id)
        add_hyperlink(p, title, url)
        r = p.add_run(" - " + note)
        set_font(r, size=10.3)

    add_heading(doc, "Internal implementation evidence supplied with this packet", 2)
    add_label_detail_table(
        doc,
        (
            ("Rule inventory", "RULE_PROVENANCE_MATRIX.md"),
            ("Claim inventory", "CLAIM_REVIEW_MATRIX.md"),
            ("Fixture inventory", "CLINICAL_REVIEW_FIXTURES.md"),
            ("Current specification", "RULES_ENGINE_SPEC.md and CURRENT_CYCLE_SUMMARY_MATRIX.md"),
            ("Implementation record", "Phase 1A and Phase 1B implementation notes; commit 261fc0f"),
            ("Automated evidence", "Core/mobile test suites, type checks, lint/build, production bundle, simulator review, and passing GitHub CI"),
        ),
    )
    add_callout(doc, "Separate counsel track", "Trademark, copyright/license, intended-use, FDA, FTC, privacy, security, App Store metadata, and public-claim substantiation are not within this clinical sign-off. Flag concerns, but do not treat this form as legal approval.")


def add_signoff(doc: Document) -> None:
    add_heading(doc, "7. Reviewer attestation and next-step record", 1)
    add_body(doc, "Sign-off applies only to the decisions and limitations explicitly recorded in this packet. Blank items remain unapproved.")
    add_label_detail_table(
        doc,
        (
            ("Reviewer name", ""),
            ("Credentials / certification", ""),
            ("Current role / organization", ""),
            ("Relevant clinical experience", ""),
            ("Authorized source edition(s)", ""),
            ("Conflict-of-interest disclosure", ""),
            ("Review date", ""),
        ),
    )

    add_heading(doc, "Attestation", 2)
    add_body(doc, "I reviewed the current software behavior, priority decision records, synthetic fixtures, terminology worksheet, and stated limitations. My decisions are limited to the scope and sources I documented. I understand that this review does not grant trademark/copyright rights, establish regulatory status, approve privacy/security claims, or create affiliation with a method owner or training organization.")

    for label in (
        "Reviewer signature:",
        "Date:",
        "Product owner acknowledgement:",
        "Engineering decision-record version:",
    ):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(label + " " + "_" * 72)
        set_font(r, size=10.5)

    add_heading(doc, "Required follow-up after review", 2)
    for text in (
        "Product converts each reviewer decision into a versioned rule/claim record with exact scope and wording.",
        "Counsel separately reviews method naming, source/license provenance, intended use, public claims, privacy, and platform metadata.",
        "Engineering updates source, specification, fixtures, tests, presentation model, and user copy together - never one in isolation.",
        "The reviewer rechecks changed gold fixtures and rendered user-facing copy before release.",
        "The release record preserves reviewer, source edition, app/engine version, limitations, and re-review trigger.",
    ):
        add_body(doc, "[ ] " + text)

    add_callout(doc, "Completion rule", "No blank, held, or disputed clinical decision may be treated as approved merely because the current automated tests pass.", PALE_GOLD)


def build() -> Path:
    doc = Document()
    configure_styles(doc)
    configure_header_footer(doc)
    bullet_id, number_id = add_numbering(doc)
    _, source_number_id = add_numbering(doc)

    props = doc.core_properties
    props.title = "Well Within Clinical Review Packet"
    props.subject = "Clinical rules, pattern boundaries, fixtures, and terminology review"
    props.author = "Well Within"
    props.keywords = "clinical review, FABM, rules engine, fixtures, Well Within"
    props.comments = "Reviewer working document; not clinical validation or legal approval."

    add_cover(doc)
    add_orientation(doc, bullet_id, number_id)
    add_current_model(doc, bullet_id)
    add_decisions(doc, bullet_id)
    add_fixtures(doc)
    add_terminology(doc)
    add_sources(doc, source_number_id)
    add_signoff(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
