from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_clinical_review_packet as base


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Well_Within_Clinical_Review_Packet_AI_Research_Draft.docx"
RESEARCH_DATE = "July 13, 2026"

SUPPORTED = "Supported"
CONTRADICTED = "Contradicted"
MIXED = "Mixed"
INSUFFICIENT = "Insufficient"

STATUS_STYLE = {
    SUPPORTED: (base.PALE_SAGE, "426A58"),
    CONTRADICTED: ("F5DEDE", "8D3E3E"),
    MIXED: (base.PALE_GOLD, "7B6427"),
    INSUFFICIENT: (base.LIGHT_GRAY, base.ACCENT_DARK),
}


@dataclass(frozen=True)
class Source:
    key: str
    citation: str
    url: str
    note: str


@dataclass(frozen=True)
class ResearchDecision:
    id: str
    title: str
    audit_ids: str
    status: str
    finding: str
    general_evidence: Sequence[tuple[str, Sequence[str]]]
    well_within: str
    draft_disposition: str
    manual_boundary: str
    reviewer_questions: Sequence[str]
    crosscheck: str = ""


SOURCES = {
    "P1": Source(
        "P1",
        "Saint Paul VI Institute. CrMS App Copyright: Samples of Intellectual Property/Copyrighted Material, p. 3 (public PDF).",
        "https://saintpaulvi.com/PDF/CrMS_App_Copyright.pdf",
        "Selected public charting concepts and sample codes. Public access is not a license or a complete software specification.",
    ),
    "P2": Source(
        "P2",
        "Creighton Model FertilityCare System. About the Creighton Model FertilityCare System.",
        "https://creightonmodel.com/about/",
        "Describes a standardized system taught with individualized follow-up by trained practitioners.",
    ),
    "P3": Source(
        "P3",
        "Creighton Model FertilityCare System. Background.",
        "https://creightonmodel.com/background/",
        "Public overview of menstrual onset, mucus progression, Peak, long-cycle mucus patches, BIP, and change from an established pattern.",
    ),
    "P4": Source(
        "P4",
        "FertilityCare Centers of America. The Creighton Model System.",
        "https://www.fertilitycare.org/creighton-model-system/",
        "Describes standardized observation and charting with practitioner instruction and follow-up.",
    ),
    "R1": Source(
        "R1",
        "Stanford JB, Smith KR, Varner MW. Impact of instruction in the Creighton Model FertilityCare System on time to pregnancy in couples of proven fecundity: results of a randomised trial. Paediatr Perinat Epidemiol. 2014;28(5):391-399. doi:10.1111/ppe.12141.",
        "https://pmc.ncbi.nlm.nih.gov/articles/PMC4861655/",
        "Research conducted with trained instruction and follow-up; it does not validate Well Within.",
    ),
    "R2": Source(
        "R2",
        "Stanford JB, Schliep KC, Chang C-P, O'Sullivan J-P, Porucznik CA. Comparison of woman-picked, expert-picked, and computer-picked Peak Day of cervical mucus with blinded urine luteinising hormone surge for concurrent identification of ovulation. Paediatr Perinat Epidemiol. 2020;34(2):105-113. doi:10.1111/ppe.12642.",
        "https://pmc.ncbi.nlm.nih.gov/articles/PMC8495767/",
        "Shows imperfect agreement among women, experts, and a different research algorithm; incomplete charts were excluded.",
    ),
    "R3": Source(
        "R3",
        "Najmabadi S, Schliep KC, Simonsen SE, Porucznik CA, Egger MJ, Stanford JB. Cervical mucus patterns and the fertile window in women without known subfertility: a pooled analysis of three cohorts. Hum Reprod. 2021;36(7):1784-1795. doi:10.1093/humrep/deab049.",
        "https://pmc.ncbi.nlm.nih.gov/articles/PMC8487651/",
        "Trained CrMS charting data showing peak-type signs, multiple mucus peaks, non-Peak-only patterns, and substantial within-woman variability.",
    ),
    "R4": Source(
        "R4",
        "Najmabadi S, Schliep KC, Simonsen SE, Porucznik CA, Egger MJ, Stanford JB. Menstrual bleeding, cycle length, and follicular and luteal phase lengths in women without known subfertility: a pooled analysis of three cohorts. Paediatr Perinat Epidemiol. 2020;34(3):318-327. doi:10.1111/ppe.12644.",
        "https://pmc.ncbi.nlm.nih.gov/articles/PMC8495765/",
        "Uses woman-identified menstrual onset and documents meaningful within-woman cycle variability.",
    ),
    "R5": Source(
        "R5",
        "Porucznik CA, Cox KJ, Schliep KC, Stanford JB. Pilot test and validation of the Peak Day method of prospective determination of ovulation against a handheld urine hormone monitor. BMC Women's Health. 2014;14:4. doi:10.1186/1472-6874-14-4.",
        "https://pmc.ncbi.nlm.nih.gov/articles/PMC3893397/",
        "Tests a simplified research Peak Day method, not the Well Within rules engine or a full CrMS software implementation.",
    ),
    "G1": Source(
        "G1",
        "Centers for Disease Control and Prevention. U.S. Medical Eligibility Criteria: Fertility Awareness-Based Methods.",
        "https://www.cdc.gov/contraception/hcp/usmec/fertility-awareness-based-methods.html",
        "Identifies circumstances in which fertility-awareness use may require delay, caution, or special counseling.",
    ),
    "B1": Source(
        "B1",
        "U.S. Food and Drug Administration. General Wellness: Policy for Low Risk Devices, Guidance for Industry and Food and Drug Administration Staff. January 2026.",
        "https://www.fda.gov/regulatory-information/search-fda-guidance-documents/general-wellness-policy-low-risk-devices",
        "Regulatory boundary source only; it does not approve this product or establish clinical correctness.",
    ),
    "B2": Source(
        "B2",
        "U.S. Food and Drug Administration. Policy for Device Software Functions and Mobile Medical Applications.",
        "https://www.fda.gov/regulatory-information/search-fda-guidance-documents/policy-device-software-functions-and-mobile-medical-applications",
        "Regulatory boundary source only; intended use and claims matter.",
    ),
    "B3": Source(
        "B3",
        "Federal Trade Commission. Health Products Compliance Guidance.",
        "https://www.ftc.gov/business-guidance/resources/health-products-compliance-guidance",
        "Advertising boundary source only; product claims require appropriately matched substantiation.",
    ),
    "B4": Source(
        "B4",
        "U.S. Copyright Office. Works Not Protected by Copyright, Circular 33.",
        "https://www.copyright.gov/circs/circ33.pdf",
        "Copyright boundary source only: methods and systems are distinct from protected expression. Counsel must assess the actual materials, license terms, and implementation.",
    ),
}


DECISIONS: tuple[ResearchDecision, ...] = (
    ResearchDecision(
        "D-01",
        "Daily observation reduction",
        "R-01; I-08",
        MIXED,
        "The general idea of recording the most fertile sign is publicly supported. The app's exact 0-3 ranking, tie handling, conflicting-entry handling, and use of a single reduced value are not publicly validated.",
        (
            (
                "The public sample instructs users to chart daily and record the most fertile sign. That supports a strongest-sign concept at the method level.",
                ("P1",),
            ),
            (
                "Research comparing woman-, expert-, and computer-selected Peak Day shows that chart reduction and interpretation are not interchangeable across people or algorithms.",
                ("R2",),
            ),
        ),
        "Insufficient. Internal tests show that Well Within applies its rank consistently and retains the original observations; no public study validates this rank, its tie rules, or its output for conflicting observations.",
        "Hold the exact software reduction rule for qualified review. A reviewer may accept the broad strongest-sign concept while separately adjudicating every mapped input and conflict case.",
        "Do not infer the proprietary ordering, exceptions, or conflict rules from a public sample. Require the reviewer to cite an authorized source edition and page or expressly document clinical judgment.",
        (
            "Is strongest-sign reduction appropriate for every supported context?",
            "For conflicting observations, should the app select, warn, require correction, or record without interpretation?",
            "Which original details must remain visible next to a reduced value?",
        ),
    ),
    ResearchDecision(
        "D-02",
        "Sensation, appearance, and generated notation",
        "R-02-R-05; R-07; I-01-I-03; I-09",
        MIXED,
        "Public material supports several observation terms and sample notation elements, but it does not validate Well Within's complete mapping, suffix order, omitted-frequency behavior, rank scale, or opening-boundary logic.",
        (
            (
                "The public sample lists selected sensation, appearance, frequency, and notation examples, including dry, damp, wet, shiny, sticky, tacky, stretchy, and several appearance letters.",
                ("P1",),
            ),
            (
                "Official descriptions place these observations inside a standardized system taught with individualized instruction and follow-up.",
                ("P2", "P4"),
            ),
        ),
        "Insufficient. No public study or official source validates the app's full observation-to-notation transformation, internal 0-3 ranks, suffix ordering, or decision to let damp, wet, or shiny open a derived pattern.",
        "Treat the public codes as orientation only. Hold the exact mapping until an authorized source-by-source review covers I-01 through I-03 and I-09.",
        "The public PDF identifies its contents as samples of copyrighted/intellectual-property material. Do not copy missing tables, ordering, or expression from an unauthorized manual, and do not treat public access as a license.",
        (
            "Which exact mappings are approved for entry, interpretation, and export?",
            "May damp, wet, or shiny without a mucus appearance open an automated boundary?",
            "Does omitted frequency affect interpretation, notation only, or entry validity?",
        ),
    ),
    ResearchDecision(
        "D-03",
        "Contradictory or incomplete input combinations",
        "R-06; I-04-I-06; I-10",
        INSUFFICIENT,
        "No authoritative public source located defines Well Within's accept, warn, reject, or record-without-interpretation policy for contradictory fields, missing required fields, duplicate dates, or invalid date order.",
        (
            (
                "The computer-comparison study excluded incomplete charts, which confirms that missingness can affect interpretability but does not establish an app input policy.",
                ("R2",),
            ),
            (
                "Official public descriptions emphasize standardized observations and practitioner follow-up; they do not publish a complete validation matrix for software input combinations.",
                ("P2", "P4"),
            ),
        ),
        "Insufficient. The current dry-plus-appearance promotions and legacy date tolerances are implementation behavior, not clinical validation.",
        "Hold automated interpretation for unresolved contradictory combinations. Ask the reviewer to define a complete minimum-valid-input and error-recovery table.",
        "A proprietary manual may contain relevant distinctions. Do not reverse-engineer them from sample codes or current app behavior.",
        (
            "For dry plus clear, lubricative, or yellow appearance, should the app accept, warn, correct, or exclude the day from interpretation?",
            "Which fields are mandatory before a day may drive a marker?",
            "How should duplicate, missing, invalid, and out-of-order dates be handled?",
        ),
    ),
    ResearchDecision(
        "D-04",
        "Flow, spotting, brown observations, and mucus",
        "R-09-R-12; F-06; I-07; C-10-C-11",
        MIXED,
        "Public evidence supports recording mucus during light or very-light flow, creating a contradiction signal for any blanket rule that makes all flow-day mucus unusable. The exact automated Peak candidacy and ordinary-user wording remain unresolved.",
        (
            (
                "The public sample specifically instructs recording the presence or absence of mucus on light and very-light flow days.",
                ("P1",),
            ),
            (
                "Peer-reviewed CrMS chart analyses include bleeding and spotting context, but do not validate Well Within's suppression, opening-boundary, or candidate rules.",
                ("R3", "R4"),
            ),
        ),
        "Insufficient for the implementation. Well Within suppresses Peak-type interpretation on menstrual-flow days and blocks light/spotting rows from opening the derived interval or becoming a candidate; no public product-specific validation was found.",
        "Do not approve a blanket flow-day exclusion from public evidence. Require fixture-level adjudication for C-10 and C-11, including what remains descriptive and what may drive a marker.",
        "The precise interaction between flow symbols, mucus signs, Peak candidacy, and later counting may require an authorized manual. Leave it unresolved until cited review.",
        (
            "Can light-flow or spotting days affect a Peak-type marker, and under what exact conditions?",
            "Which bleeding observations are descriptive only or require individualized review?",
            "What should users see when bleeding and a high-quality mucus sign occur together?",
        ),
    ),
    ResearchDecision(
        "D-05",
        "Cycle start and continuing flow",
        "R-21-R-22; F-07; C-12-C-14",
        MIXED,
        "Public evidence supports cycle onset at the beginning of menstruation, not Well Within's exact heavy/moderate-only threshold. Leading light flow, gaps within flow, and charts without heavy/moderate flow remain unresolved.",
        (
            (
                "The official public background describes the cycle as beginning with the onset of menstruation.",
                ("P3",),
            ),
            (
                "A pooled cohort study used woman-identified menstrual onset and documents substantial cycle variability; it does not validate a heavy/moderate-only software boundary.",
                ("R4",),
            ),
        ),
        "Insufficient for the implementation. Calendar-date arithmetic has been internally corrected, but the rule that starts a cycle at a non-continuing heavy/moderate row lacks public clinical validation.",
        "Keep the calendar-date correction. Hold the boundary classifier until a reviewer marks C-12 through C-14 and defines behavior for absent dates and no qualifying flow.",
        "Do not derive a bleeding threshold from the phrase onset of menstruation. Exact symbol and continuing-flow rules need authorized review.",
        (
            "For light then heavy then moderate flow, which date is Cycle Day 1?",
            "How does a missing calendar date inside heavy/moderate flow affect continuation?",
            "What should happen when no heavy/moderate boundary exists?",
        ),
    ),
    ResearchDecision(
        "D-06",
        "Peak-type observation definition",
        "R-13; C-13; I-01-I-08",
        MIXED,
        "Clear, stretchy, and lubricative signs are supported as peak-type concepts in public and peer-reviewed sources. Well Within's cloudy/clear mapping, promoted combinations, dry-plus-appearance behavior, and rank boundary are not validated.",
        (
            (
                "The public sample and official background identify clear, stretchy, or lubricative observations as peak-type signs.",
                ("P1", "P3"),
            ),
            (
                "Peer-reviewed CrMS data use clear, stretchy, or slippery/lubricative characteristics to describe peak-type mucus while also documenting more complex patterns.",
                ("R3",),
            ),
        ),
        "Insufficient for the exact implementation. No public source validates the complete set of inputs that Well Within promotes to its highest rank or labels Peak-type.",
        "A reviewer can accept the broad concept while holding every unlisted promotion and the exact set membership. Adjudicate I-01 through I-08 before release.",
        "Exact combination handling and any hierarchy beyond the public descriptors may require an authorized manual.",
        (
            "Which exact recorded observations may be called Peak-type in an unassisted app?",
            "Which combinations must never be promoted automatically?",
            "Is the proposed chart-only phrase acceptable for ordinary users?",
        ),
        "Audit cross-reference warning: the original packet cites C-13, but C-13 is the missing-date-inside-flow fixture and appears unrelated to Peak-type definition. Confirm the intended fixture.",
    ),
    ResearchDecision(
        "D-07",
        "Retrospective Peak Day and the three-day sequence",
        "R-14-R-16; C-01; C-05; C-15",
        MIXED,
        "The last clear, stretchy, or lubricative day and public P+3 framework support a retrospective Peak concept. They do not validate Well Within's strict lower-rank comparison, reset behavior, gap handling, safeguards, or wording.",
        (
            (
                "Official public sources describe Peak as the last day of clear, stretchy, or lubricative mucus; the public sample also refers to P+3.",
                ("P1", "P3"),
            ),
            (
                "A study comparing women, experts, and a different computer algorithm found imperfect agreement with each other and with an LH-based reference. Its algorithm and trained data cannot be transferred to Well Within.",
                ("R2",),
            ),
            (
                "A separate simplified Peak Day study likewise does not validate a full CrMS app rule or patient-facing claims.",
                ("R5",),
            ),
        ),
        "Insufficient for the implementation. No product-specific evidence validates the requirement that each of the next three consecutive observed days rank lower, the candidate reset, the last-seven-day safeguard, or retrospective edits.",
        "Hold the exact algorithm and wording pending authorized rule review and fixture adjudication. Keep explicit language that a chart marker does not confirm ovulation.",
        "Do not infer a full P+1/P+2/P+3 rule set from a brief public mention. Require authorized citations for exact comparisons, exceptions, and restarts.",
        (
            "Is the current three-consecutive-lower-days rule correct for the supported scope?",
            "Does gradual decline require any relationship among P+1, P+2, and P+3?",
            "Which terms are permitted, and what limitation must always appear?",
        ),
    ),
    ResearchDecision(
        "D-08",
        "Calendar gaps and days marked not observed",
        "R-25; C-02-C-04",
        INSUFFICIENT,
        "Public evidence confirms that complete, standardized observation matters, but no source located defines whether an absent date and an explicit not-observed day are equivalent or which gap should block which boundary.",
        (
            (
                "Official descriptions place daily observations inside a standardized, practitioner-supported method.",
                ("P2", "P4"),
            ),
            (
                "Research studies used trained charting and excluded incomplete charts in relevant analyses; that supports caution with missingness, not Well Within's exact blocking policy.",
                ("R1", "R2"),
            ),
        ),
        "Insufficient. The current blocking behavior is internally verified and conservative, but there is no public clinical validation for its boundary effects, recovery rules, or wording.",
        "Retain the no-conclusion posture as an interim safeguard, not as a validated clinical rule. Require C-02 through C-04 adjudication.",
        "Do not invent missing-data equivalence or recovery rules from general calls for daily observation.",
        (
            "Should an absent date and explicit not-observed row have the same effect?",
            "Which gaps limit opening, block Peak, or make the full pattern unsupported?",
            "What recovery wording is accurate without blaming the user?",
        ),
    ),
    ResearchDecision(
        "D-09",
        "Separated or repeated Peak-type sequences",
        "R-17; F-05; C-09",
        MIXED,
        "Public and peer-reviewed sources show that repeated mucus patches or multiple peaks occur and may require context. They do not establish which sequence Well Within should select or validate its nonselection wording.",
        (
            (
                "The public sample acknowledges a double-Peak situation, while the official background notes multiple mucus patches in long cycles.",
                ("P1", "P3"),
            ),
            (
                "Cohort research documents multiple mucus peaks and substantial pattern variability; expert and computer Peak selections can also differ.",
                ("R3", "R2"),
            ),
        ),
        "Insufficient for the implementation. Showing a review state and selecting neither sequence is a conservative product choice, not a validated CrMS classification.",
        "Keep nonselection as an interim safety posture. A qualified reviewer must decide whether any sequence can be selected and under what complete prerequisites.",
        "Double-Peak and long-cycle adjudication may rely on proprietary rules and individualized instruction. Do not infer those rules.",
        (
            "Is nonselection plus optional review the correct default for two qualifying sequences?",
            "May the app ever select the later sequence, and under what exact prerequisites?",
            "What wording explains nonselection without diagnosis or implied certainty?",
        ),
    ),
    ResearchDecision(
        "D-10",
        "Continuous mucus, non-Peak-only patterns, and special contexts",
        "R-19-R-20; C-06-C-08; special-context set",
        MIXED,
        "Public sources support that BIP, continuous discharge, non-Peak-only patterns, and physiologic contexts can require different handling. They do not validate a single automated interpretation path for Well Within.",
        (
            (
                "The official background describes an established Basic Infertile Pattern and change from that pattern with proper instruction.",
                ("P3",),
            ),
            (
                "Research shows some women have only non-Peak-type mucus and that continuous fluid may be uninformative for a simplified Peak algorithm.",
                ("R3", "R5"),
            ),
            (
                "CDC guidance identifies postpartum/breastfeeding, perimenopause, irregular bleeding, discharge, and some drug-related circumstances as requiring delay, caution, or special counseling for fertility-awareness methods.",
                ("G1",),
            ),
        ),
        "Insufficient for implementation-specific interpretation. Well Within currently withholds a confirmed summary and has not invented BIP or special-context mechanics; no public validation establishes the exact state, message, population, or prerequisites.",
        "Maintain no-automation or defined caution states until a reviewer defines supported population, minimum history, BIP prerequisites, and referral triggers.",
        "Do not encode BIP, postpartum, perimenopause, medication, infection, or subfertility rules without authorized sources and appropriate clinical review.",
        (
            "For each special context, choose full, restricted, caution-only, or no interpretation.",
            "What established history or practitioner-defined BIP is required?",
            "What neutral message and referral path should appear when interpretation is withheld?",
        ),
    ),
    ResearchDecision(
        "D-11",
        "Permitted ordinary-user terminology",
        "C-01-C-19; R-18; R-31-R-32",
        MIXED,
        "Chart-based Peak and mucus-to-P+3 terminology are supported at a method level, but imperfect agreement, additional method rules, and context dependence argue against ovulation certainty or a complete fertile-window claim. Well Within's exact phrases have not been product-validated or usability-tested.",
        (
            (
                "Public method sources use Peak Day and P+3 terminology within taught charting.",
                ("P1", "P3"),
            ),
            (
                "The comparison study shows that a chart-selected Peak Day is not identical to expert, computer, or hormone-based identification in every cycle.",
                ("R2",),
            ),
        ),
        "Insufficient for the exact implementation. No public evidence validates 'possible Peak Day,' 'post-Peak pattern,' 'Possible fertile pattern,' or the proposed limitations as understood by ordinary Well Within users.",
        "Continue holding Fertile Start/End, Total fertile days, past/closed-window, confirmed ovulation, confidence, predictive, diagnostic, and avoidance-guidance language. Phase 1C may test the exact qualified phrase 'Possible fertile pattern — based on your logged observations' only for the defined eligible scope and with the adjacent limitation.",
        "Permitted method terminology, trademarks, copyrighted expression, regulatory classification, and advertising claims are separate questions. Clinical acceptance alone cannot resolve them.",
        (
            "Approve, revise, or prohibit each exact ordinary-user phrase, including the Phase 1C possible-pattern copy.",
            "What chart-only limitation must appear next to retrospective Peak statements?",
            "Which phrases could imply ovulation, diagnosis, prediction, or avoiding-pregnancy guidance?",
        ),
    ),
    ResearchDecision(
        "D-12",
        "History comparisons and inclusion eligibility",
        "R-29-R-30; C-21-C-22",
        INSUFFICIENT,
        "Studies document substantial cycle and mucus variability, but no authoritative public evidence supports Well Within's current comparison bands, labels, minimum sample size, or eligible-cycle rules.",
        (
            (
                "Pooled studies document within-woman variability in mucus patterns, cycle length, and phase lengths.",
                ("R3", "R4"),
            ),
            (
                "Those research distributions do not validate product labels such as usual, consistent, or significant variation, nor the current numeric thresholds.",
                ("R3", "R4"),
            ),
        ),
        "Insufficient. The current plus/minus bands and inclusion rules are product thresholds. Internal date corrections and support-state filtering improve consistency but do not establish clinical meaning.",
        "Keep active-cycle predictions suppressed. Phase 1C may show raw Cycle Day ranges with N across eligible completed cycles, without usual/expected labels, normative bands, or projection onto the active/future cycle.",
        "Do not convert cohort variability into a patient-specific clinical threshold without a defined validation question and product-specific evidence.",
        (
            "Which comparisons are purely descriptive and suitable for ordinary users?",
            "What minimum eligible sample size and exclusions are required?",
            "Should bands be removed, shown numerically without labels, or validated against defined evidence?",
        ),
    ),
    ResearchDecision(
        "D-13",
        "Possible fertile pattern presentation",
        "R-33-R-35; C-25-C-27; PFP-01-PFP-06",
        MIXED,
        "Public evidence supports a beginning-of-mucus-through-P+3 concept within taught CrMS, but that is not the complete method fertility classification and does not validate Well Within's exact date band, eligibility, or wording.",
        (
            (
                "The public Saint Paul VI sample includes beginning of mucus through three full days past Peak among days of fertility, but also includes bleeding, isolated/non-Peak mucus, unusual bleeding, other counts, and a trained-instruction warning. Official background also describes long-cycle mucus patches, BIP, and change-from-pattern complexity.",
                ("P1", "P3"),
            ),
        ),
        "Insufficient for the exact implementation. Well Within has not validated its opening set, Peak/P+3 behavior, immediate response to later signs, special-context eligibility, wording comprehension, or decision safety.",
        "For the working first release, show only an unbounded 'may be developing' statement while forming. Show exact opening-through-P+3 dates only for an eligible summary-available chart, suppress them for blocked/review/unsupported charts, and keep historical ranges retrospective with N. Treat this as a product rule-conformance target, not clinical validation.",
        "Do not infer the public sample's additional bleeding, single-day, non-Peak, uncertainty, BIP, or special-context mechanics. Public access does not provide a license to reproduce protected expression or a complete software specification.",
        (
            "Is the exact eligible summary copy acceptable: 'Based on the observations you entered, this chart shows a possible fertile pattern from [start] through P+3 [end]'?",
            "Confirm that forming charts show no exact dates/band and that later signs immediately reopen, withhold, or route an earlier bounded result to review.",
            "Confirm PFP-01 through PFP-06, including context suppression, retrospective history ranges, and the adjacent limitation.",
        ),
    ),
)


REVIEWER_QUESTIONS = (
    "Define the supported population and pattern scope, including exclusions, no-interpretation states, and referral triggers.",
    "Using a named authorized manual edition and page citations, approve the complete observation-to-notation-to-classification mapping and strongest-sign reduction, including I-01 through I-10.",
    "Define minimum valid input and accept/warn/reject/record-without-interpretation behavior for contradictions, missing frequency, and date problems.",
    "Decide Cycle Day 1 and bleeding-plus-mucus treatment for C-10 through C-14: leading light flow, light/very-light flow with mucus, spotting/brown, gaps in flow, and no heavy/moderate boundary.",
    "Specify the exact Peak-type, candidate, Peak Day, P+1 through P+3, reset, gradual-decline, gap, not-observed, and past-edit rules for the first supported release; confirm that P+3 is not presented as ovulation confirmation.",
    "Specify treatment of repeated sequences, continuous mucus, non-Peak-only patterns, BIP, and special contexts, including every no-automation or referral state.",
    "Approve or prohibit exact ordinary-user terminology and adjacent limitations for in-app and exported views, including the Phase 1C Possible fertile pattern copy and PFP-01 through PFP-06; separately route regulatory, legal, trademark, copyright, and licensing questions.",
    "Define history eligibility, minimum sample size, exclusions, and whether numeric comparisons or labels are permitted; then independently adjudicate an authorized, de-identified conformance set.",
)


def set_cell_text(cell, text: str, *, size: float = 9.0, bold: bool = False, color: str = base.INK) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    base.set_font(r, size=size, bold=bold, color=color)


def add_small_body(doc: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        base.set_font(r, size=9.4, bold=True)
        r = p.add_run(text[len(bold_lead):])
        base.set_font(r, size=9.4)
    else:
        r = p.add_run(text)
        base.set_font(r, size=9.4)


def add_status_banner(doc: Document, status: str, finding: str) -> None:
    fill, color = STATUS_STYLE[status]
    table = doc.add_table(rows=1, cols=2)
    base.set_table_geometry(table, [1700, 7660])
    base.set_table_borders(table, color=fill, size="2")
    base.set_row_cant_split(table.rows[0])
    left, right = table.rows[0].cells
    base.set_cell_shading(left, fill)
    base.set_cell_shading(right, fill)
    set_cell_text(left, f"EVIDENCE\n{status.upper()}", size=9.2, bold=True, color=color)
    set_cell_text(right, finding, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_citation(doc: Document, key: str) -> None:
    src = SOURCES[key]
    p = doc.add_paragraph(style="Small Note")
    p.paragraph_format.left_indent = Inches(0.17)
    p.paragraph_format.first_line_indent = Inches(-0.17)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.02
    r = p.add_run(f"[{key}] ")
    base.set_font(r, size=7.8, color=base.ACCENT_DARK, bold=True)
    base.add_hyperlink(p, src.citation, src.url)


def add_evidence_item(doc: Document, statement: str, source_keys: Sequence[str], bullet_id: int) -> None:
    p = doc.add_paragraph()
    base.apply_num(p, bullet_id)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.06
    r = p.add_run(statement)
    base.set_font(r, size=9.0)
    for key in source_keys:
        add_citation(doc, key)


def add_reviewer_block(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=1)
    base.set_table_geometry(table, [base.PAGE_WIDTH_DXA])
    base.set_table_borders(table, color=base.BORDER)
    base.set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    base.set_cell_shading(cell, base.PALE_SAGE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Qualified reviewer disposition — intentionally blank")
    base.set_font(r, size=9.0, color=base.ACCENT_DARK, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("[ ] Accept for defined scope   [ ] Revise   [ ] No automated interpretation   [ ] Hold for authorized evidence")
    base.set_font(r, size=8.4)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Scope, required change, source edition/page, or notes:  " + "_" * 67)
    base.set_font(r, size=8.2, color=base.MUTED)


def configure_document(doc: Document) -> tuple[int, int]:
    base.configure_styles(doc)
    styles = doc.styles
    styles["Normal"].font.size = Pt(9.4)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.10
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 2"].font.size = Pt(11.5)
    styles["Heading 2"].paragraph_format.space_before = Pt(8)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    styles["Small Note"].font.size = Pt(7.8)
    styles["Small Note"].paragraph_format.space_after = Pt(2)
    doc.core_properties.title = "Well Within Clinical Review Packet - AI Research Draft"
    doc.core_properties.subject = "Public-evidence prefill for qualified clinical review"
    doc.core_properties.author = "Well Within - AI Research Draft"
    doc.core_properties.comments = (
        "Research working document only. Not clinical, legal, regulatory, trademark, licensing, or product approval."
    )
    return base.add_numbering(doc)


def configure_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("WELL WITHIN  |  AI RESEARCH DRAFT  |  CLINICAL DECISION RECORDS")
    base.set_font(r, size=8.2, color=base.MUTED, bold=True)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    base.set_table_geometry(table, [7650, 1710], indent_dxa=0)
    left, right = table.rows[0].cells
    left_p = left.paragraphs[0]
    left_p.paragraph_format.space_after = Pt(0)
    r = left_p.add_run(f"AI Research Draft | Not approval | {RESEARCH_DATE}")
    base.set_font(r, size=7.7, color=base.MUTED)
    base.add_page_field(right.paragraphs[0])

    first_footer = section.first_page_footer
    p = first_footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AI Research Draft | Public-evidence prefill only | No reviewer attestation or approval")
    base.set_font(r, size=7.8, color=base.MUTED)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("AI RESEARCH DRAFT")
    base.set_font(r, size=12, color=base.ACCENT, bold=True)

    title = doc.add_paragraph("Well Within", style="Title")
    base.set_paragraph_keep(title, keep_next=True)
    subtitle = doc.add_paragraph("Clinical Decision Records | Public-Evidence Prefill", style="Subtitle")
    base.set_paragraph_keep(subtitle, keep_next=True)

    meta = doc.add_table(rows=4, cols=2)
    rows = (
        ("Prepared for", "A qualified current Creighton/FABM clinical reviewer"),
        ("Purpose", "Reduce review time by separating public support, product-specific unknowns, and required decisions"),
        ("Evidence cutoff", RESEARCH_DATE),
        ("Implementation context", "Action 5 audit documents; no app or rules-engine changes"),
    )
    for row, (label, value) in zip(meta.rows, rows):
        base.set_row_cant_split(row)
        base.set_cell_shading(row.cells[0], base.LIGHT_GRAY)
        set_cell_text(row.cells[0], label, size=9.2, bold=True, color=base.ACCENT_DARK)
        set_cell_text(row.cells[1], value, size=9.2)
    base.set_table_geometry(meta, [1900, 7460])
    base.set_table_borders(meta)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    base.add_callout(
        doc,
        "Research conclusion",
        "No public source located validates Well Within's complete implementation. Ten decisions have mixed evidence: the general method concept is supported, but the exact software rule is unresolved. Three decisions have insufficient evidence. The product owner reports a practitioner working direction, but this AI draft does not create credentials, attestation, signature, or product-specific validation.",
        base.PALE_GOLD,
    )
    base.add_callout(
        doc,
        "Use boundary",
        "This editable draft is research support, not clinical, legal, regulatory, trademark, copyright, licensing, or product approval. It does not complete reviewer credentials, attestation, or signature. It does not infer rules from an unauthorized Creighton manual.",
        base.PALE_ROSE,
    )
    base.add_callout(
        doc,
        "Reviewer path",
        "Start with the action dashboard and eight cross-cutting questions. Then review only the decision pages and fixtures relevant to the initial supported release.",
        base.PALE_SAGE,
    )
    base.add_callout(
        doc,
        "Phase 1C working direction",
        "For planning, treat the interim clinical recommendations as accepted and add a narrowly eligible 'Possible fertile pattern.' While forming, show no exact dates or band. For an eligible retrospective chart, the proposed band runs from the first approved mucus observation through P+3. This records a product-owner-supplied assumption, not a signed reviewer approval.",
        base.PALE_SAGE,
    )


def add_dashboard(doc: Document, bullet_id: int) -> None:
    doc.add_page_break()
    base.add_heading(doc, "Action dashboard", 1)
    add_small_body(
        doc,
        "Fast read: Public evidence supports several broad Creighton concepts, but none of the 13 exact Well Within decisions has product-specific clinical validation. Internal tests establish consistency, not correctness. Phase 1C records a working clinical direction without creating an attestation.",
    )

    counts = ((MIXED, 10, "General concept partly supported; exact software rule still needs review."), (INSUFFICIENT, 3, "No adequate public evidence for the decision."), (SUPPORTED, 0, "No decision is fully supported for Well Within."), (CONTRADICTED, 0, "No whole decision is contradicted; D-04 and D-05 contain contradiction signals."))
    table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ("Rating", "Count", "Meaning in this draft")):
        base.set_cell_shading(cell, base.PALE_ROSE)
        set_cell_text(cell, value, size=8.8, bold=True, color=base.ACCENT_DARK)
    for status, count, meaning in counts:
        row = table.add_row()
        fill, color = STATUS_STYLE[status]
        base.set_cell_shading(row.cells[0], fill)
        set_cell_text(row.cells[0], status, size=8.8, bold=True, color=color)
        set_cell_text(row.cells[1], str(count), size=8.8, bold=True)
        set_cell_text(row.cells[2], meaning, size=8.8)
        base.set_row_cant_split(row)
    base.set_table_geometry(table, [1800, 900, 6660])
    base.set_table_borders(table)

    base.add_heading(doc, "Decision map", 2)
    map_table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(map_table.rows[0].cells, ("Decision", "Evidence", "Next qualified-review action")):
        base.set_cell_shading(cell, base.PALE_ROSE)
        set_cell_text(cell, value, size=8.3, bold=True, color=base.ACCENT_DARK)
    next_actions = {
        "D-01": "Adjudicate ranking and conflicts",
        "D-02": "Approve full mapping table",
        "D-03": "Define input-validity matrix",
        "D-04": "Decide bleeding + mucus treatment",
        "D-05": "Define Cycle Day 1 boundary",
        "D-06": "Approve exact Peak-type set",
        "D-07": "Approve exact retrospective rule",
        "D-08": "Define missing-data effects",
        "D-09": "Define repeated-sequence handling",
        "D-10": "Set population/context scope",
        "D-11": "Approve exact user language",
        "D-12": "Define or remove comparisons",
        "D-13": "Confirm exact eligible display",
    }
    for decision in DECISIONS:
        row = map_table.add_row()
        fill, color = STATUS_STYLE[decision.status]
        set_cell_text(row.cells[0], f"{decision.id}  {decision.title}", size=7.9, bold=True)
        base.set_cell_shading(row.cells[1], fill)
        set_cell_text(row.cells[1], decision.status, size=7.9, bold=True, color=color)
        set_cell_text(row.cells[2], next_actions[decision.id], size=7.9)
        base.set_row_cant_split(row)
    base.set_table_geometry(map_table, [4350, 1550, 3460])
    base.set_table_borders(map_table)

    base.add_heading(doc, "Immediate sequencing", 2)
    for item in (
        "First: define supported population, special-context exclusions, and no-interpretation states.",
        "Second: adjudicate observation mapping, input validity, bleeding, cycle start, Peak-type, and Peak rules.",
        "Third: align the Phase 1C possible-pattern presentation, user terminology, and history displays only after the underlying rules are fixed.",
        "Parallel outside this clinical review: counsel assesses license, copyright, trademark, regulatory, and claim boundaries.",
    ):
        base.add_bullet(doc, item, bullet_id)


def add_evidence_method(doc: Document, bullet_id: int) -> None:
    doc.add_page_break()
    base.add_heading(doc, "How to read the evidence ratings", 1)
    definitions = (
        (SUPPORTED, "Authoritative public evidence directly supports the decision as stated, including the Well Within implementation."),
        (CONTRADICTED, "Authoritative public evidence conflicts with the decision as stated."),
        (MIXED, "Some concept-level support exists, but the exact implementation is unresolved, context-dependent, or contains contrary signals."),
        (INSUFFICIENT, "Available public evidence does not answer the decision with adequate specificity."),
    )
    table = doc.add_table(rows=0, cols=2)
    for status, meaning in definitions:
        row = table.add_row()
        fill, color = STATUS_STYLE[status]
        base.set_cell_shading(row.cells[0], fill)
        set_cell_text(row.cells[0], status, size=8.9, bold=True, color=color)
        set_cell_text(row.cells[1], meaning, size=8.9)
        base.set_row_cant_split(row)
    base.set_table_geometry(table, [1800, 7560])
    base.set_table_borders(table)

    base.add_heading(doc, "Two evidence layers kept separate", 2)
    for item in (
        "Creighton/general evidence: official public descriptions and peer-reviewed studies of trained charting support or qualify a broad concept.",
        "Well Within-specific validation: evidence that the app's exact inputs, transformation, algorithm, missing-data behavior, wording, and population work as intended. No such validation was located.",
        "Implementation context: Action 5 audits and tests describe what the product does and whether it is internally consistent. They are not clinical evidence.",
    ):
        base.add_bullet(doc, item, bullet_id)

    base.add_heading(doc, "Research boundaries", 2)
    for item in (
        "Public material was not used to reconstruct missing proprietary rules.",
        "A public sample that identifies copyrighted material was treated as evidence about selected published concepts, not as a license or complete specification.",
        "Clinical evidence was not treated as legal, regulatory, trademark, copyright, licensing, or advertising approval.",
        "Questions remain unresolved wherever the public record is incomplete or does not match the exact product rule.",
        "The requested 20-chart set should be called a reviewer conformance/adjudication set, not clinical validation, unless a qualified protocol defines a validation study.",
    ):
        base.add_bullet(doc, item, bullet_id)

    base.add_callout(
        doc,
        "Implementation context reviewed",
        "ACTION_5_GOAL, REPORT, CORRECTNESS_FINDINGS, RULE_PROVENANCE_MATRIX, SOURCE_REGISTER, CLAIM_REVIEW_MATRIX, CLINICAL_REVIEW_FIXTURES, UX_SURFACING_SPEC, PHASE_1A_IMPLEMENTATION, and PHASE_1B_IMPLEMENTATION.",
        base.LIGHT_GRAY,
    )


def add_decision_page(doc: Document, decision: ResearchDecision, bullet_id: int, first: bool) -> None:
    if not first:
        doc.add_page_break()
    base.add_heading(doc, f"{decision.id}  |  {decision.title}", 1)
    add_status_banner(doc, decision.status, decision.finding)

    summary = doc.add_table(rows=2, cols=2)
    base.set_table_geometry(summary, [1750, 7610])
    base.set_table_borders(summary)
    for row, (label, value) in zip(
        summary.rows,
        (
            ("Audit references", decision.audit_ids),
            ("Current behavior", next(item.current for item in base.DECISIONS if item.id == decision.id)),
        ),
    ):
        base.set_cell_shading(row.cells[0], base.LIGHT_GRAY)
        set_cell_text(row.cells[0], label, size=8.3, bold=True, color=base.ACCENT_DARK)
        set_cell_text(row.cells[1], value, size=8.3)
        base.set_row_cant_split(row)

    base.add_heading(doc, "Public evidence about Creighton / general FABM", 2)
    for statement, source_keys in decision.general_evidence:
        add_evidence_item(doc, statement, source_keys, bullet_id)

    base.add_heading(doc, "Evidence validating Well Within's specific implementation", 2)
    base.add_callout(doc, "PRODUCT-SPECIFIC EVIDENCE", decision.well_within, base.LIGHT_GRAY)

    base.add_heading(doc, "AI draft disposition for qualified review", 2)
    add_small_body(doc, decision.draft_disposition)
    add_small_body(doc, "Authorized-source boundary: " + decision.manual_boundary, bold_lead="Authorized-source boundary: ")
    if decision.crosscheck:
        base.add_callout(doc, "CROSS-REFERENCE CHECK", decision.crosscheck, base.PALE_GOLD)

    base.add_heading(doc, "Questions the reviewer must decide", 2)
    for question in decision.reviewer_questions:
        base.add_bullet(doc, question, bullet_id)
    add_reviewer_block(doc)


def add_reviewer_questions(doc: Document, number_id: int) -> None:
    doc.add_page_break()
    base.add_heading(doc, "Eight decisions to formalize in the qualified-review record", 1)
    add_small_body(
        doc,
        "These are the smallest cross-cutting decisions needed to turn the 13 records into an implementable clinical specification. A response should state the supported scope and cite an authorized source edition/page where method-specific rules are used.",
    )
    for question in REVIEWER_QUESTIONS:
        base.add_numbered(doc, question, number_id)

    base.add_heading(doc, "What this review can and cannot close", 2)
    base.add_callout(
        doc,
        "Clinical reviewer can close",
        "Supported population, observation mappings, rule behavior, missing-data treatment, fixture outcomes, clinical limitations, and suitable chart-based terminology.",
        base.PALE_SAGE,
    )
    base.add_callout(
        doc,
        "Separate specialist review remains",
        "License, copyright, trademark, affiliation, FDA/device status, advertising claims, privacy, and platform policy. A clinical answer must not be reused as approval in those areas.",
        base.PALE_GOLD,
    )
    add_small_body(
        doc,
        "Recommended evidence artifact: an authorized, de-identified chart conformance set independently marked by the reviewer, with discrepancies adjudicated against a versioned decision record. Do not describe a small conformance set as clinical validation without a qualified study protocol.",
    )


def add_source_register(doc: Document) -> None:
    doc.add_page_break()
    base.add_heading(doc, "Public evidence register", 1)
    add_small_body(
        doc,
        "All sources below were publicly accessible on the evidence cutoff date. Links are editable and clickable. Method-level evidence does not validate Well Within, and boundary sources do not constitute approval.",
    )
    for key in ("P1", "P2", "P3", "P4", "R1", "R2", "R3", "R4", "R5", "G1", "B1", "B2", "B3", "B4"):
        src = SOURCES[key]
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.04
        r = p.add_run(f"[{src.key}] ")
        base.set_font(r, size=8.4, color=base.ACCENT_DARK, bold=True)
        base.add_hyperlink(p, src.citation, src.url)
        p2 = doc.add_paragraph(style="Small Note")
        p2.paragraph_format.left_indent = Inches(0.22)
        p2.paragraph_format.space_after = Pt(4)
        r = p2.add_run(src.note)
        base.set_font(r, size=7.8, color=base.MUTED)

    base.add_heading(doc, "Internal implementation context", 2)
    add_small_body(
        doc,
        "The Action 5 audit documents were used only to describe current behavior, known contradictions, fixture coverage, UI suppression, and test consistency. They are repository records, not external clinical evidence.",
    )
    p = doc.add_paragraph(style="Small Note")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.02
    r = p.add_run(
        "Reviewed: REPORT, CORRECTNESS_FINDINGS, RULE_PROVENANCE_MATRIX, SOURCE_REGISTER, "
        "CLAIM_REVIEW_MATRIX, CLINICAL_REVIEW_FIXTURES, UX_SURFACING_SPEC, "
        "PHASE_1A_IMPLEMENTATION, and PHASE_1B_IMPLEMENTATION."
    )
    base.set_font(r, size=7.6)


def add_closing(doc: Document) -> None:
    doc.add_page_break()
    base.add_heading(doc, "Draft status and handoff", 1)
    base.add_callout(
        doc,
        "Current status",
        "AI research prefill and the product-owner-supplied practitioner working direction are recorded. A formal qualified-review record is still required for the exact rules, fixtures, scope, and wording. No reviewer credentials, attestation, or signature has been completed here.",
        base.PALE_GOLD,
    )
    base.add_heading(doc, "Completion checklist", 2)
    for item in (
        "[ ] Supported population and exclusions defined",
        "[ ] Authorized source edition/pages supplied for method-specific rules",
        "[ ] All 13 decision records adjudicated",
        "[ ] Fixtures relevant to the first release independently marked",
        "[ ] User-facing terminology approved with limitations",
        "[ ] History comparison policy defined or removed",
        "[ ] PFP-01 through PFP-06 and exact Possible fertile pattern copy adjudicated",
        "[ ] Separate legal/regulatory/trademark/copyright/licensing review completed",
        "[ ] Versioned product-specific validation plan defined before any validation claim",
    ):
        add_small_body(doc, item)
    base.add_callout(
        doc,
        "No product change",
        "This research draft does not modify the Well Within app, rules engine, tests, terminology, exports, or release state.",
        base.PALE_SAGE,
    )


def build() -> Path:
    doc = Document()
    bullet_id, number_id = configure_document(doc)
    configure_header_footer(doc)
    add_cover(doc)
    add_dashboard(doc, bullet_id)
    add_evidence_method(doc, bullet_id)

    doc.add_page_break()
    base.add_heading(doc, "Clinical decision records", 1)
    add_small_body(
        doc,
        "Each record separates broad method evidence from evidence for the exact Well Within implementation. The reviewer disposition is intentionally blank.",
    )
    for index, decision in enumerate(DECISIONS):
        add_decision_page(doc, decision, bullet_id, first=(index == 0))

    add_reviewer_questions(doc, number_id)
    add_source_register(doc)
    add_closing(doc)

    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.72)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
